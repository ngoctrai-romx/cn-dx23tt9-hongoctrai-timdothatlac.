from zipfile import ZipFile, ZIP_DEFLATED
from pathlib import Path
from xml.sax.saxutils import escape
from pptx import Presentation
from pptx.util import Inches, Pt

OUT = Path('do_an_nop_qlht')
OUT.mkdir(exist_ok=True)
student_id='170123090'
name='Hồ Ngọc Trai'
klass='DX23TT9'
teacher='Nguyễn Nhứt Lam'
title='XÂY DỰNG WEBSITE TÌM KIẾM ĐỒ THẤT LẠC'

def p(text, bold=False, center=False, size=24):
    jc = '<w:jc w:val="center"/>' if center else ''
    b = '<w:b/>' if bold else ''
    return f'''<w:p><w:pPr>{jc}</w:pPr><w:r><w:rPr>{b}<w:sz w:val="{size}"/><w:szCs w:val="{size}"/></w:rPr><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'''

def heading(text, level=1):
    size = 32 if level==1 else 28
    return p(text, bold=True, size=size)

paras=[]
paras += [p('TRƯỜNG ĐẠI HỌC TRÀ VINH', True, True, 28), p('KHOA KỸ THUẬT VÀ CÔNG NGHỆ', True, True, 26), p('', size=24), p('BÁO CÁO ĐỒ ÁN CHUYÊN NGÀNH', True, True, 36), p(f'Đề tài: {title}', True, True, 32), p('', size=24), p(f'Sinh viên thực hiện: {name}', False, True, 26), p(f'Mã số sinh viên: {student_id}', False, True, 26), p(f'Lớp: {klass}', False, True, 26), p(f'Giảng viên hướng dẫn: {teacher}', False, True, 26), p('', size=24), p('Trà Vinh, năm 2026', False, True, 24)]
sections = [
('LỜI MỞ ĐẦU','Trong đời sống hằng ngày, việc thất lạc các vật dụng cá nhân như giấy tờ tùy thân, ví tiền, điện thoại, chìa khóa, laptop hoặc thú cưng diễn ra khá phổ biến. Người bị mất đồ thường đăng tin rời rạc trên mạng xã hội, trong khi người nhặt được đồ cũng khó tìm đúng chủ sở hữu. Vì vậy, đề tài “Xây dựng website tìm kiếm đồ thất lạc” được thực hiện nhằm tạo ra một nền tảng trực tuyến hỗ trợ đăng tin, tìm kiếm và kết nối cộng đồng.'),
('CHƯƠNG 1: GIỚI THIỆU ĐỀ TÀI','Đề tài tập trung xây dựng website Tìm Đồ Thất Lạc, cho phép người dùng đăng tin mất đồ hoặc nhặt được đồ, tìm kiếm bài đăng theo từ khóa, lọc theo danh mục và chọn vị trí trên bản đồ. Hệ thống hướng đến giao diện thân thiện, dễ sử dụng và có khả năng mở rộng trong tương lai.'),
('1.1. Lý do chọn đề tài','Thông tin mất đồ hiện nay thường bị phân tán trên nhiều nền tảng khác nhau như Facebook, Zalo hoặc thông báo thủ công. Các bài đăng dễ bị trôi, khó tìm kiếm và thiếu bộ lọc. Website chuyên biệt giúp tập trung dữ liệu, hỗ trợ người dùng tìm kiếm nhanh hơn và tăng khả năng hoàn trả đồ vật.'),
('1.2. Mục tiêu đề tài','Xây dựng website có các chức năng: đăng ký, đăng nhập, đăng tin mất đồ/nhặt được đồ, xem danh sách bài đăng, tìm kiếm, lọc theo loại tin và danh mục, chọn vị trí trên bản đồ và lưu dữ liệu vào MongoDB.'),
('CHƯƠNG 2: PHÂN TÍCH YÊU CẦU','Người dùng có thể xem trang chủ, tìm kiếm đồ vật, xem danh sách bài đăng, đăng ký, đăng nhập và đăng tin mới. Khi đăng tin, người dùng nhập tiêu đề, danh mục, thời gian, mô tả, hình ảnh minh họa và chọn vị trí trên bản đồ.'),
('2.1. Yêu cầu chức năng','Các chức năng chính gồm: quản lý tài khoản người dùng, xác thực đăng nhập, quản lý bài đăng, tìm kiếm bài đăng theo tiêu đề hoặc mô tả, lọc theo loại tin lost/found, lọc theo danh mục và tích hợp bản đồ chọn vị trí.'),
('2.2. Yêu cầu phi chức năng','Giao diện hiện đại, dễ sử dụng, có khả năng hiển thị trên nhiều thiết bị. Mật khẩu người dùng được mã hóa bằng bcryptjs. Dữ liệu được lưu trong MongoDB thông qua Mongoose. Hệ thống có thể mở rộng thêm các chức năng nâng cao.'),
('CHƯƠNG 3: THIẾT KẾ HỆ THỐNG','Hệ thống được xây dựng bằng Next.js App Router. Frontend sử dụng React để tạo giao diện. Backend sử dụng API Routes của Next.js. Database sử dụng MongoDB. Xác thực người dùng dùng NextAuth. Bản đồ sử dụng Leaflet và OpenStreetMap.'),
('3.1. Kiến trúc hệ thống','Người dùng tương tác với giao diện website. Các thao tác như đăng ký, đăng nhập, lấy danh sách bài đăng và tạo bài đăng được gửi đến API Route. API xử lý nghiệp vụ, xác thực session và lưu hoặc đọc dữ liệu từ MongoDB.'),
('3.2. Thiết kế dữ liệu','Collection User gồm các trường name, phone, email, password, createdAt, updatedAt. Collection Post gồm title, type, category, categoryName, description, date, location, image, author, authorName, status, createdAt, updatedAt.'),
('CHƯƠNG 4: CÔNG NGHỆ SỬ DỤNG','Dự án sử dụng Next.js 16, React 19, MongoDB, Mongoose, NextAuth, bcryptjs, Leaflet, React-Leaflet và OpenStreetMap. Các công nghệ này phù hợp để xây dựng ứng dụng web hiện đại, có frontend và backend trong cùng một project.'),
('CHƯƠNG 5: XÂY DỰNG CHỨC NĂNG','Trang chủ gồm phần giới thiệu, thanh tìm kiếm, hành động nhanh, danh mục và bài đăng gần đây. Trang bài đăng hỗ trợ xem danh sách, tìm kiếm và lọc. Trang đăng tin có form nhập liệu và bản đồ chọn tọa độ. Hệ thống API xử lý lấy danh sách bài đăng và tạo bài đăng mới.'),
('5.1. Chức năng đăng ký và đăng nhập','Người dùng đăng ký bằng họ tên, số điện thoại, email và mật khẩu. Email được kiểm tra trùng lặp. Mật khẩu được mã hóa bằng bcryptjs trước khi lưu. Đăng nhập sử dụng NextAuth để quản lý phiên làm việc.'),
('5.2. Chức năng đăng tin','Người dùng đã đăng nhập có thể đăng tin mới. Nếu chưa đăng nhập, hệ thống yêu cầu đăng nhập. Khi tạo bài đăng, dữ liệu được lưu vào MongoDB, kèm thông tin người đăng và trạng thái active.'),
('5.3. Chức năng bản đồ','Component MapPicker sử dụng Leaflet để hiển thị bản đồ. Người dùng có thể tìm địa điểm bằng OpenStreetMap Nominatim, click vào bản đồ để chọn tọa độ và xóa vị trí để chọn lại.'),
('CHƯƠNG 6: KẾT QUẢ ĐẠT ĐƯỢC','Dự án đã hoàn thành giao diện trang chủ, trang danh sách bài đăng, trang đăng tin, đăng ký, đăng nhập, API quản lý bài đăng, kết nối MongoDB, mã hóa mật khẩu và tích hợp bản đồ chọn vị trí. Project có thể build production thành công.'),
('CHƯƠNG 7: HẠN CHẾ VÀ HƯỚNG PHÁT TRIỂN','Hạn chế: chức năng upload ảnh thật chưa hoàn thiện, chưa có trang quản trị, chưa có chat trực tiếp, chưa xác thực email/số điện thoại. Hướng phát triển: thêm upload Cloudinary/Firebase, admin dashboard, chat, thông báo, tìm kiếm theo khoảng cách và ứng dụng di động.'),
('KẾT LUẬN','Đề tài đã xây dựng được website tìm kiếm đồ thất lạc đáp ứng các chức năng cơ bản. Qua quá trình thực hiện, sinh viên vận dụng kiến thức về lập trình web, React/Next.js, API, MongoDB, xác thực người dùng và tích hợp bản đồ vào ứng dụng thực tế.'),
('TÀI LIỆU THAM KHẢO','Next.js Documentation; React Documentation; MongoDB Documentation; Mongoose Documentation; NextAuth Documentation; Leaflet Documentation; OpenStreetMap.')]
for h,t in sections:
    paras.append(heading(h))
    for part in t.split('\n'):
        paras.append(p(part, size=24))

doc_xml = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>{''.join(paras)}<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr></w:body></w:document>'''
ct='''<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'''
rels='''<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'''
with ZipFile(OUT/f'{student_id}_{name}_{klass}_BaoCao.docx','w',ZIP_DEFLATED) as z:
    z.writestr('[Content_Types].xml',ct); z.writestr('_rels/.rels',rels); z.writestr('word/document.xml',doc_xml)

prs=Presentation()
slides=[('ĐỒ ÁN CHUYÊN NGÀNH', title+'\nSinh viên: '+name+'\nMSSV: '+student_id+'\nGVHD: '+teacher),('Lý do chọn đề tài','Thông tin mất đồ phân tán, khó tìm kiếm; cần nền tảng tập trung để kết nối người mất và người nhặt được.'),('Mục tiêu','Đăng ký/đăng nhập\nĐăng tin mất đồ/nhặt được đồ\nTìm kiếm, lọc bài đăng\nChọn vị trí trên bản đồ\nLưu dữ liệu MongoDB'),('Công nghệ sử dụng','Next.js, React, MongoDB, Mongoose, NextAuth, bcryptjs, Leaflet, OpenStreetMap'),('Chức năng chính','Trang chủ\nDanh sách bài đăng\nĐăng tin mới\nBản đồ chọn vị trí\nAPI quản lý bài đăng'),('Thiết kế dữ liệu','User: name, phone, email, password\nPost: title, type, category, description, location, author, status'),('Kết quả đạt được','Hoàn thiện giao diện và chức năng cơ bản\nCó xác thực người dùng\nCó lưu dữ liệu MongoDB\nCó tích hợp bản đồ'),('Hạn chế & phát triển','Hoàn thiện upload ảnh\nThêm admin dashboard\nThêm chat/thông báo\nTìm kiếm theo khoảng cách'),('Video thuyết trình','Link video thuyết trình: [Dán link Google Drive/YouTube tại đây]'),('Kết luận','Website đáp ứng mục tiêu đề tài và có thể phát triển thành sản phẩm hỗ trợ cộng đồng tìm lại đồ thất lạc.')]
for i,(t,b) in enumerate(slides):
    slide=prs.slides.add_slide(prs.slide_layouts[0] if i==0 else prs.slide_layouts[1])
    slide.shapes.title.text=t
    body=slide.placeholders[1]
    body.text=b
    for paragraph in body.text_frame.paragraphs:
        paragraph.font.size=Pt(24)
prs.save(OUT/f'{student_id}_{name}_{klass}_ThuyetTrinh.pptx')

script = f'''KỊCH BẢN THUYẾT TRÌNH VIDEO\n\nKính chào thầy {teacher}. Em tên là {name}, MSSV {student_id}, lớp {klass}. Hôm nay em xin trình bày đồ án chuyên ngành với đề tài: {title}.\n\nLý do em chọn đề tài là vì trong thực tế việc thất lạc đồ vật cá nhân diễn ra khá thường xuyên. Người mất đồ thường đăng tin trên mạng xã hội nhưng thông tin dễ bị trôi, còn người nhặt được đồ cũng khó tìm đúng chủ sở hữu. Vì vậy em xây dựng website hỗ trợ đăng tin và tìm kiếm đồ thất lạc.\n\nWebsite có các chức năng chính gồm đăng ký, đăng nhập, đăng tin mất đồ hoặc nhặt được đồ, xem danh sách bài đăng, tìm kiếm theo từ khóa, lọc theo danh mục và chọn vị trí trên bản đồ.\n\nVề công nghệ, em sử dụng Next.js và React để xây dựng giao diện, API Routes để xử lý backend, MongoDB và Mongoose để lưu dữ liệu, NextAuth để xác thực người dùng, bcryptjs để mã hóa mật khẩu và Leaflet kết hợp OpenStreetMap để hiển thị bản đồ.\n\nKết quả đạt được là website đã có giao diện trang chủ, trang danh sách bài đăng, trang đăng tin, chức năng đăng ký đăng nhập, kết nối cơ sở dữ liệu MongoDB và tích hợp bản đồ chọn tọa độ.\n\nDo thời gian còn hạn chế, hệ thống vẫn cần phát triển thêm chức năng upload ảnh thật, trang quản trị, chat giữa người dùng và thông báo tự động.\n\nEm xin chân thành cảm ơn thầy đã theo dõi phần trình bày của em.'''
(OUT/f'{student_id}_{name}_{klass}_KichBanThuyetTrinh.txt').write_text(script, encoding='utf-8')
print('created', OUT.resolve())
